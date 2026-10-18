"""Kia: total road tax + other charges (no insurance, no TCS) -> FinStack Logic expression.

Source: Desktop\\kia_before_insurance_correction.xlsx, Sheet1.
Value = total_road_tax_inr (base tax + 11% cess + Rs1000 fixed cess) + registration_other_allowance_inr.
Key   = snake_case of Car_Type (model + variant), same rule as kia_prices_snake.json.
"""
import json
from pathlib import Path

import openpyxl

SRC = Path(r"C:\Users\samsm\OneDrive\Desktop\kia_before_insurance_correction.xlsx")
OUT = Path(__file__).resolve().parent


def key(name):
    return "_".join("".join(c if c.isalnum() else " " for c in name.replace("+", " plus ").lower()).split())


ws = openpyxl.load_workbook(SRC, data_only=True).active
head = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(head)}
out, rows = {}, 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[col["vehicle_id"]] is None:
        continue
    rows += 1
    # check the sheet's own parts add up before trusting the total
    tax = r[col["base_road_tax_inr"]] + r[col["cess_on_tax_inr"]] + r[col["fixed_welfare_cess_inr"]]
    assert tax == r[col["total_road_tax_inr"]], (r[col["vehicle_id"]], tax, r[col["total_road_tax_inr"]])
    k = key(r[col["Car_Type"]])
    assert k not in out, f"duplicate key {k}"
    out[k] = int(r[col["total_road_tax_inr"]] + r[col["registration_other_allowance_inr"]])

(OUT / "kia_roadtax_others_snake.json").write_text(json.dumps(out, indent=4), encoding="utf-8")
body = json.dumps(out, separators=(",", ":"))   # one line, as short as possible to copy
(OUT / "kia_roadtax_others_expression.txt").write_text(f"{body} if {{#common/lead/count#}} is not None else {body}", encoding="utf-8")
print(rows, "rows,", len(out), "keys; min", min(out.values()), "max", max(out.values()))
